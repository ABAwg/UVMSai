import io
import json
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

import jwt
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.main import app
from app.config import settings
from app.database import SessionLocal
from app.models import RefreshSession, Project, MemoryItem, Subscription
from app.security import decode_token


def register(client, web=False):
    response = client.post('/api/v1/auth/register', json={'email': f'qa-{uuid.uuid4().hex}@example.com', 'password': 'QA-regression-123', 'display_name': 'QA'}, headers={'X-UVMS-Client': 'web'} if web else {})
    assert response.status_code == 201, response.text
    pair = response.json()
    return pair, {'Authorization': f"Bearer {pair['access_token']}"}


def test_expired_access_cookie_recovery_and_logout():
    with TestClient(app) as client:
        pair, headers = register(client, web=True)
        assert pair['refresh_token'] is None
        cookie = client.cookies.get('uvms_refresh')
        claims = decode_token(pair['access_token'], 'access')
        claims['exp'] = int((datetime.now(timezone.utc) - timedelta(seconds=2)).timestamp())
        expired = jwt.encode(claims, settings.jwt_secret, algorithm='HS256')
        assert client.get('/api/v1/me', headers={'Authorization': f'Bearer {expired}'}).status_code == 401
        refresh = client.post('/api/v1/auth/refresh', json={}, headers={'X-UVMS-Client': 'web'})
        assert refresh.status_code == 200
        assert refresh.json()['refresh_token'] is None
        assert 'HttpOnly' in refresh.headers['set-cookie']
        assert client.cookies.get('uvms_refresh') != cookie
        headers = {'Authorization': f"Bearer {refresh.json()['access_token']}"}
        assert client.get('/api/v1/me', headers=headers).status_code == 200
        assert client.post('/api/v1/auth/logout', json={}).status_code == 204
        assert client.get('/api/v1/me', headers=headers).status_code == 401
        assert client.post('/api/v1/auth/refresh', json={}).status_code == 401


def test_expired_refresh_and_untrusted_origin():
    with TestClient(app) as client:
        pair, headers = register(client)
        assert client.post('/api/v1/auth/refresh', json={}, headers={'Origin': 'https://untrusted.example'}).status_code == 403
        with SessionLocal() as db:
            row = db.scalar(select(RefreshSession).where(RefreshSession.jti == decode_token(pair['refresh_token'], 'refresh')['jti']))
            row.expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
            db.commit()
        response = client.post('/api/v1/auth/refresh', json={})
        assert response.status_code == 401
        assert response.json()['error']['code'] == 'http_401'
        assert not client.cookies.get('uvms_refresh')


def test_parallel_rotation_accepts_one_refresh_only():
    with TestClient(app) as client:
        pair, _ = register(client)
        def rotate(_):
            with TestClient(app) as other:
                return other.post('/api/v1/auth/refresh', json={'refresh_token': pair['refresh_token']}).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            assert sorted(pool.map(rotate, range(2))) == [200, 401]


def test_logout_all_invalidates_access_on_every_device():
    with TestClient(app) as client:
        pair, headers = register(client)
        other = client.post('/api/v1/auth/login', json={'email': pair['user']['email'], 'password': 'QA-regression-123'}).json()
        assert client.post('/api/v1/me/logout-all', headers=headers).status_code == 204
        for access in [pair['access_token'], other['access_token']]:
            assert client.get('/api/v1/me', headers={'Authorization': f'Bearer {access}'}).status_code == 401


def test_rotation_preserves_other_tab_access_and_stable_device_session():
    with TestClient(app) as client:
        pair, headers = register(client)
        first = client.get('/api/v1/me/sessions', headers=headers).json()
        refreshed = client.post('/api/v1/auth/refresh', json={})
        assert refreshed.status_code == 200
        assert client.get('/api/v1/me', headers=headers).status_code == 200
        second = client.get('/api/v1/me/sessions', headers=headers).json()
        assert [row['id'] for row in first] == [row['id'] for row in second]
        assert client.post('/api/v1/auth/refresh', json={'refresh_token': pair['refresh_token']}).status_code == 401


def test_search_failure_falls_back_without_claiming_live_sources(monkeypatch):
    from app.ai_gateway import stream_events
    def provider(*args, web_search=False, **kwargs):
        if web_search:
            raise RuntimeError('search temporarily unavailable')
        yield {'type':'delta','delta':'Answer without web sources'}
    monkeypatch.setattr('app.ai_gateway._provider_stream_events', provider)
    events = list(stream_events([], 'uvms-1-balanced', web_search=True))
    assert events[0] == {'type':'tool','tool':'web_search','status':'failed'}
    assert events[1]['delta'] == 'Answer without web sources'


@pytest.mark.parametrize('action', ['edit', 'regenerate'])
def test_failed_replacement_preserves_history(monkeypatch, action):
    with TestClient(app) as client:
        _, headers = register(client)
        result = client.post('/api/v1/chat/complete', headers=headers, json={'content': 'Original question'}).json()
        path = f"/api/v1/chats/{result['chat_id']}"
        before = client.get(path, headers=headers).json()['messages']
        def fail(*args, **kwargs):
            raise RuntimeError('provider unavailable')
        monkeypatch.setattr('app.routes.complete_answer', fail)
        message = before[0 if action == 'edit' else 1]
        response = client.post(f"{path}/messages/{message['id']}/{action}", headers=headers, json={'content': 'Edited question'} if action == 'edit' else {})
        assert response.status_code == 502
        assert client.get(path, headers=headers).json()['messages'] == before


@pytest.mark.parametrize('partial', [False, True])
def test_failed_stream_is_retryable_without_duplicate_prompt(monkeypatch, partial):
    with TestClient(app) as client:
        _, headers = register(client)
        def fail(*args, **kwargs):
            if partial:
                yield {'type': 'delta', 'delta': 'Unfinished answer'}
            raise RuntimeError('provider unavailable')
        monkeypatch.setattr('app.routes.stream_events', fail)
        first = client.post('/api/v1/chat/stream', headers=headers, json={'content': 'Retry this'})
        assert 'event: error' in first.text
        chat_id = first.headers['x-uvms-chat-id']
        message_id = first.headers['x-uvms-message-id']
        messages = client.get(f'/api/v1/chats/{chat_id}', headers=headers).json()['messages']
        assert len(messages) == 1 and messages[0]['role'] == 'user'
        monkeypatch.setattr('app.routes.stream_events', lambda *args, **kwargs: iter([{'type': 'delta', 'delta': 'Recovered'}]))
        retry = client.post('/api/v1/chat/stream', headers=headers, json={'chat_id': chat_id, 'retry_message_id': message_id})
        assert 'event: done' in retry.text
        messages = client.get(f'/api/v1/chats/{chat_id}', headers=headers).json()['messages']
        assert [m['content'] for m in messages] == ['Retry this', 'Recovered']
        assert client.post('/api/v1/chat/stream', headers=headers, json={'chat_id': chat_id, 'retry_message_id': message_id}).status_code == 409


def test_cancelled_stream_keeps_only_nonempty_marked_partial(monkeypatch):
    with TestClient(app) as client:
        _, headers = register(client)
        calls = 0
        def cancelled(_):
            nonlocal calls
            calls += 1
            return calls > 1
        monkeypatch.setattr('app.routes.is_cancelled', cancelled)
        monkeypatch.setattr('app.routes.stream_events', lambda *args, **kwargs: iter([{'type': 'delta', 'delta': 'First'}, {'type': 'delta', 'delta': 'Second'}]))
        response = client.post('/api/v1/chat/stream', headers=headers, json={'content': 'Stop me'})
        messages = client.get('/api/v1/chats/' + response.headers['x-uvms-chat-id'], headers=headers).json()['messages']
        assert messages[-1]['content'] == 'First'
        assert messages[-1]['meta']['stopped'] is True


def test_project_files_instructions_and_enabled_memory_reach_provider(monkeypatch):
    captured = []
    def answer(history, *args, **kwargs):
        captured.extend(history)
        return 'Context received', {}
    monkeypatch.setattr('app.routes.complete_answer_with_meta', answer)
    with TestClient(app) as client:
        _, headers = register(client)
        project = client.post('/api/v1/projects', headers=headers, json={'name': 'Context', 'instructions': 'Use metric units.'}).json()
        client.post('/api/v1/files', headers=headers, data={'project_id': project['id']}, files={'upload': ('notes.txt', b'Project launch date: Friday', 'text/plain')})
        client.post('/api/v1/memory', headers=headers, json={'content': 'I prefer concise answers.'})
        disabled = client.post('/api/v1/memory', headers=headers, json={'content': 'Do not include this memory.'}).json()
        client.patch('/api/v1/memory/' + disabled['id'], headers=headers, json={'is_enabled': False})
        response = client.post('/api/v1/chat/complete', headers=headers, json={'content': 'What is the plan?', 'project_id': project['id']})
        assert response.status_code == 200, response.text
        context = json.dumps(captured)
        assert 'Use metric units.' in context and 'Project launch date: Friday' in context
        assert 'I prefer concise answers.' in context and 'Do not include this memory.' not in context
        assert client.get('/api/v1/chats/' + response.json()['chat_id'], headers=headers).json()['project']['id'] == project['id']


def test_account_delete_cascades_sqlite_data():
    with TestClient(app) as client:
        pair, headers = register(client)
        client.post('/api/v1/projects', headers=headers, json={'name': 'Delete QA'})
        client.post('/api/v1/memory', headers=headers, json={'content': 'Disposable test data'})
        assert client.delete('/api/v1/me', headers=headers).status_code == 204
        with SessionLocal() as db:
            for model in [Project, MemoryItem, Subscription, RefreshSession]:
                assert db.scalar(select(model).where(model.user_id == pair['user']['id'])) is None


def test_file_formats_and_validation():
    from docx import Document
    from openpyxl import Workbook
    from pypdf import PdfWriter
    doc = Document(); doc.add_paragraph('Document QA'); docx = io.BytesIO(); doc.save(docx)
    book = Workbook(); book.active['A1'] = 'Spreadsheet QA'; xlsx = io.BytesIO(); book.save(xlsx)
    writer = PdfWriter(); writer.add_blank_page(width=200, height=200); pdf = io.BytesIO(); writer.write(pdf)
    with TestClient(app) as client:
        _, headers = register(client)
        for name, data in [('qa.docx', docx.getvalue()), ('qa.xlsx', xlsx.getvalue()), ('qa.pdf', pdf.getvalue()), ('qa.txt', b'Text QA')]:
            upload = client.post('/api/v1/files', headers=headers, files={'upload': (name, data)})
            assert upload.status_code == 201, upload.text
            assert client.get('/api/v1/files/' + upload.json()['id'] + '/download', headers=headers).content == data
        assert client.post('/api/v1/files', headers=headers, files={'upload': ('qa.exe', b'unsupported')}).status_code == 415


def test_model_restrictions_and_demo_copy():
    with TestClient(app) as client:
        _, headers = register(client)
        assert client.post('/api/v1/chat/complete', headers=headers, json={'content': 'Hello', 'model_id': 'uvms-1-deep'}).status_code == 403
        response = client.post('/api/v1/chat/complete', headers=headers, json={'content': 'Hello'})
        assert 'demo' in response.json()['answer'].lower()
        assert 'OPENAI_API_KEY' not in response.text and '.env' not in response.text
        assert response.json()['meta']['demo'] is True


def test_history_search_count_pagination_and_owner_isolation():
    from app.models import Chat, Message
    with TestClient(app) as client:
        pair, headers = register(client)
        other, _ = register(client)
        with SessionLocal() as db:
            chats = [Chat(user_id=pair['user']['id'], title=f'History {i}') for i in range(103)]
            db.add_all(chats)
            private = Chat(user_id=other['user']['id'], title='Private needle')
            db.add(private)
            db.flush()
            db.add(Message(chat_id=chats[0].id, role='user', content='needle 100% literal_value'))
            db.commit()
        response = client.get('/api/v1/chats', headers=headers)
        assert len(response.json()) == 100 and response.headers['X-Total-Count'] == '103'
        page = client.get('/api/v1/chats?offset=100', headers=headers)
        assert len(page.json()) == 3
        assert not {row['id'] for row in page.json()} & {row['id'] for row in response.json()}
        for query in ['needle', '100%', 'literal_value']:
            found = client.get('/api/v1/chats', params={'q': query}, headers=headers)
            assert found.headers['X-Total-Count'] == '1'
            assert found.json()[0]['title'] == 'History 0'


def test_rate_limit_identity_survives_rotation_and_rejects_forgery():
    from app.rate_limit import identity_key, _memory_check
    from app.security import create_access_token
    uid = uuid.uuid4().hex
    first = identity_key('127.0.0.1', 'Bearer ' + create_access_token(uid))
    second = identity_key('127.0.0.1', 'Bearer ' + create_access_token(uid))
    assert first == second == f'user:{uid}'
    assert _memory_check(first, 1, 60).allowed
    assert not _memory_check(second, 1, 60).allowed
    assert identity_key('127.0.0.1', 'Bearer forged-token') == 'ip:127.0.0.1'


def test_unexpected_errors_use_consistent_safe_schema(monkeypatch):
    import app.main as main
    def broken_readiness():
        raise RuntimeError('Private diagnostic detail')
    monkeypatch.setattr(main, 'readiness', broken_readiness)
    with TestClient(app) as client:
        response = client.get('/ready')
        assert response.status_code == 500
        assert response.json()['error']['code'] == 'internal_error'
        assert 'Private diagnostic detail' not in response.text
        assert response.headers['X-Request-ID']
