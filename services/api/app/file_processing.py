from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import Path, PurePosixPath

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .config import settings

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css", ".sql", ".yaml", ".yml", ".xml", ".log"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx"}
ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS


def safe_filename(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._()\- ]+", "_", Path(name).name).strip()
    return cleaned[:180] or "file"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def category_for(filename: str, mime_type: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXTENSIONS or mime_type.startswith("image/"):
        return "image"
    if ext in DOCUMENT_EXTENSIONS:
        return "document"
    return "text"


def _validate_zip_document(data: bytes, ext: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if len(names) > 5000:
                raise ValueError("Office archive has too many entries")
            total_uncompressed = 0
            for info in archive.infolist():
                path = PurePosixPath(info.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ValueError("Unsafe Office archive path")
                total_uncompressed += max(0, info.file_size)
                if total_uncompressed > 150 * 1024 * 1024:
                    raise ValueError("Office archive expands beyond safety limit")
            required = "word/document.xml" if ext == ".docx" else "xl/workbook.xml"
            if required not in names:
                raise ValueError("File extension does not match Office document content")
    except zipfile.BadZipFile as exc:
        raise ValueError("Invalid Office document") from exc


def validate_upload(filename: str, data: bytes, mime_type: str) -> None:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file extension")
    if not data:
        raise ValueError("Empty file")
    if ext in IMAGE_EXTENSIONS:
        from .vision import validate_image
        validate_image(data)
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("Invalid PDF signature")
    if ext == ".png" and not data.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ValueError("Invalid PNG signature")
    if ext in {".jpg", ".jpeg"} and not data.startswith(b"\xff\xd8\xff"):
        raise ValueError("Invalid JPEG signature")
    if ext == ".gif" and not data.startswith((b"GIF87a", b"GIF89a")):
        raise ValueError("Invalid GIF signature")
    if ext == ".webp" and not (len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP"):
        raise ValueError("Invalid WEBP signature")
    if ext in {".docx", ".xlsx"}:
        _validate_zip_document(data, ext)
    # Text is allowed to contain replacement characters, but reject obvious NUL-heavy binary payloads.
    if ext in TEXT_EXTENSIONS and data.count(b"\x00") > max(8, len(data) // 100):
        raise ValueError("Text file appears to contain binary data")


def extract_text(filename: str, data: bytes, max_chars: int | None = None) -> str | None:
    max_chars = max_chars or settings.max_extracted_text_chars
    ext = Path(filename).suffix.lower()
    try:
        if ext in TEXT_EXTENSIONS:
            text = data.decode("utf-8", errors="replace")
        elif ext == ".pdf":
            reader = PdfReader(io.BytesIO(data), strict=False)
            chunks: list[str] = []
            for page in reader.pages[:80]:
                chunks.append(page.extract_text() or "")
                if sum(len(x) for x in chunks) >= max_chars:
                    break
            text = "\n\n".join(chunks)
        elif ext == ".docx":
            doc = Document(io.BytesIO(data))
            chunks = []
            total = 0
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    chunks.append(paragraph.text)
                    total += len(paragraph.text)
                    if total >= max_chars:
                        break
            text = "\n".join(chunks)
        elif ext == ".xlsx":
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            chunks = []
            total = 0
            for ws in wb.worksheets[:12]:
                chunks.append(f"# Sheet: {ws.title}")
                for idx, row in enumerate(ws.iter_rows(values_only=True)):
                    if idx >= 500 or total >= max_chars:
                        chunks.append("[sheet truncated]")
                        break
                    vals = ["" if value is None else str(value) for value in row[:40]]
                    if any(vals):
                        line = "\t".join(vals)
                        chunks.append(line)
                        total += len(line)
            text = "\n".join(chunks)
        else:
            return None
    except Exception as exc:
        return f"[Text extraction failed: {type(exc).__name__}]"

    text = text.strip()
    if not text:
        return "[No extractable text found]"
    if len(text) > max_chars:
        return text[:max_chars] + "\n[content truncated]"
    return text
